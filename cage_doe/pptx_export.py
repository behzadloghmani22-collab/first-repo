"""Build a PowerPoint deck from the campaign figures (python-pptx, optional)."""
from __future__ import annotations

import logging
from pathlib import Path

import numpy as np

log = logging.getLogger(__name__)

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Inches, Pt
    HAVE_PPTX = True
except ImportError:  # pragma: no cover
    HAVE_PPTX = False

INK = (0x0B, 0x0B, 0x0B)
INK2 = (0x52, 0x51, 0x4E)
ACCENT = (0x2A, 0x78, 0xD6)


def _image_size(path: Path) -> tuple[int, int]:
    try:
        from PIL import Image
        with Image.open(path) as im:
            return im.size
    except Exception:  # noqa: BLE001
        return (1600, 900)


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.W, self.H = self.prs.slide_width, self.prs.slide_height
        self.blank = self.prs.slide_layouts[6]

    def _title(self, slide, text, sub=None):
        tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), self.W - Inches(1.0), Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size, p.font.bold = Pt(26), True
        p.font.color.rgb = RGBColor(*INK)
        if sub:
            p2 = tf.add_paragraph()
            p2.text = sub
            p2.font.size = Pt(13)
            p2.font.color.rgb = RGBColor(*INK2)

    def _fit_picture(self, slide, path: Path, left, top, width, height):
        w, h = _image_size(path)
        scale = min(width / w, height / h)
        pw, ph = int(w * scale), int(h * scale)
        slide.shapes.add_picture(str(path), left + (width - pw) // 2, top + (height - ph) // 2, pw, ph)

    def _bullets(self, slide, items, left, top, width, height, size=14):
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame
        tf.word_wrap = True
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = "• " + it
            p.font.size = Pt(size)
            p.font.color.rgb = RGBColor(*INK)
            p.space_after = Pt(6)

    def title_slide(self, title, subtitle, note=None):
        s = self.prs.slides.add_slide(self.blank)
        tb = s.shapes.add_textbox(Inches(0.8), Inches(2.3), self.W - Inches(1.6), Inches(1.5))
        p = tb.text_frame.paragraphs[0]
        p.text = title
        p.font.size, p.font.bold = Pt(36), True
        p.font.color.rgb = RGBColor(*INK)
        p2 = tb.text_frame.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(18)
        p2.font.color.rgb = RGBColor(*INK2)
        if note:
            tb2 = s.shapes.add_textbox(Inches(0.8), Inches(5.6), self.W - Inches(1.6), Inches(1))
            q = tb2.text_frame.paragraphs[0]
            q.text = note
            q.font.size = Pt(14)
            q.font.color.rgb = RGBColor(0xEB, 0x68, 0x34)

    def figure_slide(self, title, path: Path | None, sub=None, bullets=None):
        s = self.prs.slides.add_slide(self.blank)
        self._title(s, title, sub)
        top = Inches(1.25)
        if bullets and path is not None:
            self._bullets(s, bullets, Inches(0.5), top, Inches(4.0), self.H - top - Inches(0.4))
            self._fit_picture(s, path, Inches(4.6), top, self.W - Inches(5.0), self.H - top - Inches(0.4))
        elif path is not None:
            self._fit_picture(s, path, Inches(0.5), top, self.W - Inches(1.0), self.H - top - Inches(0.4))
        elif bullets:
            self._bullets(s, bullets, Inches(0.7), top, self.W - Inches(1.4), self.H - top - Inches(0.5), size=18)

    def two_figure_slide(self, title, left_path, right_path, sub=None):
        s = self.prs.slides.add_slide(self.blank)
        self._title(s, title, sub)
        top = Inches(1.25)
        half = (self.W - Inches(1.2)) // 2
        if left_path is not None:
            self._fit_picture(s, left_path, Inches(0.5), top, half, self.H - top - Inches(0.4))
        if right_path is not None:
            self._fit_picture(s, right_path, Inches(0.7) + half, top, half, self.H - top - Inches(0.4))

    def table_slide(self, title, df, sub=None, fmt="{:.3g}"):
        s = self.prs.slides.add_slide(self.blank)
        self._title(s, title, sub)
        rows, cols = df.shape
        shape = s.shapes.add_table(rows + 1, cols, Inches(0.5), Inches(1.4), self.W - Inches(1.0),
                                   Inches(0.4) * (rows + 1))
        t = shape.table
        for j, c in enumerate(df.columns):
            t.cell(0, j).text = str(c)
            t.cell(0, j).text_frame.paragraphs[0].font.size = Pt(11)
            t.cell(0, j).text_frame.paragraphs[0].font.bold = True
        for i in range(rows):
            for j, c in enumerate(df.columns):
                v = df.iloc[i, j]
                txt = fmt.format(v) if isinstance(v, (float, np.floating)) and np.isfinite(v) else \
                    ("" if isinstance(v, (float, np.floating)) else str(v))
                t.cell(i + 1, j).text = txt
                t.cell(i + 1, j).text_frame.paragraphs[0].font.size = Pt(10)

    def save(self, path: Path):
        self.prs.save(str(path))


def build_deck(cr, figures: dict[str, Path], out: Path | None = None) -> Path | None:
    if not HAVE_PPTX:
        log.warning("python-pptx not installed - skipping the PowerPoint deck (pip install python-pptx)")
        return None
    cfg = cr.cfg
    out = Path(out or (cr.results_dir / "presentation" / "lattice_cage_DOE_results.pptx"))
    out.parent.mkdir(parents=True, exist_ok=True)
    dk = Deck()
    names = ", ".join(vr.variant.name for vr in cr.variants)
    dk.title_slide(cfg["project"]["name"], f"Design of experiments and optimisation of {len(cr.variants)} lattice variants "
                   f"({names})\nnTopCL campaign, {cr.started[:10]}",
                   "SURROGATE DEMO - figures generated from the mock model, not from nTop" if cr.mock else None)
    o, ph = cfg["objectives"], cfg["physics"]
    dk.figure_slide("What 'optimal' means", figures.get("objective_definition"),
                    "Weighted geometric mean of clinical desirabilities; any hard-limit violation scores zero",
                    [f"Apparent modulus {o['modulus_band_GPa'][0]}-{o['modulus_band_GPa'][1]} GPa (bone/PEEK-like) - weight {o['weights']['modulus']:.2f}",
                     f"Porosity {o['porosity_band'][0]:.0%}-{o['porosity_band'][1]:.0%} - weight {o['weights']['porosity']:.2f}",
                     f"Pore size {o['pore_band_um'][0]:.0f}-{o['pore_band_um'][1]:.0f} um - weight {o['weights']['pore_size']:.2f}",
                     f"Yield SF >= {o['sf_yield_min']} at {ph['peak_load_N']:.0f} N (hard); fatigue SF >= {o['sf_fatigue_min']} at {ph['design_load_N']:.0f} N",
                     f"Strut >= {o['min_strut_mm']} mm (LPBF); mass and surface area as tie-breakers"])
    doe = cfg["doe"]
    dk.figure_slide("Method: sequential DOE driven through nTopCL", figures.get("run_summary"),
                    "Every run is one ntopcl.exe call with a JSON input file; outputs are parsed automatically",
                    [", ".join(f"{f['name']} {f['low']}-{f['high']} {f['units']}" for f in cfg["factors"]),
                     f"Screening: {doe.get('screening')} + {doe.get('center_points')} centre points",
                     f"Response surface: {doe.get('rsm')}; {doe.get('lhs_runs')} Latin-hypercube runs",
                     f"Adaptive infill: {doe.get('infill_iterations')} x {doe.get('infill_per_iteration')} runs (exploit / explore / diversify)",
                     "Surrogates: quadratic RSM vs thin-plate RBF, chosen by leave-one-out error",
                     f"Verification: top {doe.get('verify_top_n')} predicted designs re-run in nTop"])
    dk.figure_slide("Ranking of lattice types", figures.get("ranking"))
    dk.figure_slide("Clinical metrics of the recommended designs", figures.get("compare_metrics"))
    dk.two_figure_slide("Desirability profile and ranking robustness", figures.get("radar"), figures.get("weight_robustness"))
    dk.figure_slide("Trade-off fronts of all lattice types", figures.get("pareto_all"),
                    "stiffness matching vs fatigue strength, porosity vs fatigue strength (feasible designs only)")
    rk = cr.ranking
    if not rk.empty:
        cols = ["rank", "variant", "D_recommended"] + [f["name"] for f in cfg["factors"] if f["name"] in rk] + \
               ["E_app_GPa", "porosity", "pore_um", "sf_yield", "sf_fatigue", "mass_g"]
        dk.table_slide("Recommended design per lattice", rk[[c for c in cols if c in rk]])
    dk.figure_slide("Recommended factor settings", figures.get("optimum_settings"))
    if "predicted_vs_verified" in figures:
        dk.figure_slide("Surrogate vs nTop verification", figures["predicted_vs_verified"])
    for vr in cr.variants:
        n = vr.variant.name
        rec = vr.recommended or {}
        sub = (", ".join(f"{f} = {rec[f]:.3g}" for f in vr.names if np.isfinite(rec.get(f, np.nan))) +
               f"  |  D = {rec.get('D', float('nan')):.3f}") if rec else "no feasible design found"
        dk.figure_slide(f"{n}: response surfaces", figures.get(f"{n}/response_surfaces"), sub)
        dk.two_figure_slide(f"{n}: main effects and sensitivity", figures.get(f"{n}/main_effects"), figures.get(f"{n}/sensitivity"), sub)
        dk.two_figure_slide(f"{n}: trade-off and interactions", figures.get(f"{n}/pareto"), figures.get(f"{n}/interactions"), sub)
    if not rk.empty:
        best = rk.iloc[0]
        dk.figure_slide("Recommendation", None, None, [
            f"{best['variant']} ({best['lattice']}) ranks first with D = {best['D_recommended']:.3f} ({best['source']})",
            "Settings: " + ", ".join(f"{f['name']} = {best[f['name']]:.3g} {f['units']}" for f in cfg["factors"]
                                    if f["name"] in rk and np.isfinite(best[f["name"]])),
            f"E_app = {best['E_app_GPa']:.2f} GPa, porosity = {best['porosity']:.2f}, pore = {best['pore_um']:.0f} um, "
            f"SF yield = {best['sf_yield']:.1f}, SF fatigue = {best['sf_fatigue']:.1f}, mass = {best['mass_g']:.1f} g",
            "Ranking robustness: see the weight-perturbation slide - a robust winner keeps rank 1 in most perturbations",
            "Next steps: confirm with a refined FE mesh and an ASTM F2077 compression / fatigue test of the printed design"])
    for vr in cr.variants:
        n = vr.variant.name
        dk.two_figure_slide(f"Appendix - {n}: sampled designs and convergence", figures.get(f"{n}/design_space"),
                            figures.get(f"{n}/convergence"))
        if f"{n}/model_fit" in figures:
            dk.figure_slide(f"Appendix - {n}: surrogate fit quality", figures[f"{n}/model_fit"])
    dk.save(out)
    return out
